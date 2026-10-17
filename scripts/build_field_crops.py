"""Add the source-verified Rice-to-Niger cohort without rebuilding accepted crops.

Usage: python scripts/build_field_crops.py SOURCE_DIRECTORY
Required snapshots: Cross Walks-2.0.docx, production-context.docx,
statewide-production.xlsx, district-apy.xlsx, originals/, catalog-manifest.json,
drive-assets.json and editor-notes.json. No network access is used by this build.
"""
import hashlib
import json
import re
import sys
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as E
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1])
LIB = ROOT / 'src/lib/gramshakti'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
CROSSWALK = 'https://docs.google.com/document/d/1J3ZQGbI51KarrX-dpLAqXTtFyjptqtav/edit?usp=drivesdk&ouid=111951905071183633266&rtpof=true&sd=true'
CONTEXT = 'https://drive.google.com/file/d/1e9EXVD5b2MfbSr7UyhieBx8YooXf-wP1/view'
STATE = 'https://docs.google.com/spreadsheets/d/1JkNGc4fvKO9_dXegJKpTuj5TZSxNWt4CaAnSJsn5oa8/edit'
DISTRICT = 'https://docs.google.com/spreadsheets/d/1YQLRV8QGIk9lO1MygZA9heofVSnxtZpDHSkm2yzTYLg/edit'
SECTIONS = {
    'rice': ('Rice', 2888, 3017), 'wheat': ('Wheat', 3017, 3140),
    'jowar': ('Jowar', 3263, 3387), 'bajra': ('Bajra', 3387, 3505),
    'barley': ('Barley', 3505, 3641), 'ragi': ('Ragi', 3641, 3756),
    'small-millets': ('Small Millets', 4005, 4154), 'gram': ('Gram', 4154, 4283),
    'pigeon-pea': ('Pigeon Pea', 4283, 4398), 'black-gram': ('Black Gram', 4398, 4514),
    'moong': ('Moong', 4514, 4631), 'lentil': ('Lentil', 4631, 4752),
    'cotton': ('Cotton', 4752, 4885), 'jute-mesta': ('Jute & Mesta', 4885, 5078),
    'tobacco': ('Tobacco', 5078, 5265), 'guarseed': ('Guarseed', 5265, 5455),
    'groundnut': ('Groundnut', 5455, 5659), 'rapeseed-mustard': ('Rapeseed & Mustard', 5659, 5874),
    'soybean': ('Soybean', 5874, 6069), 'sunflower': ('Sunflower', 6069, 6257),
    'sesamum': ('Sesamum', 6257, 6478), 'safflower': ('Safflower', 6478, 6678),
    'castor': ('Castor', 6678, 6880), 'niger-seeds': ('Niger Seeds', 6880, 10000),
}
def key(value): return re.sub(r'[^a-z0-9]', '', str(value or '').lower().replace('&', 'and'))
def read(name): return json.loads((LIB / name).read_text())
def write(name, value): (LIB / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')
codes = read('state-codes.json')
aliases = {key(n): c for n, c in codes.items()}
aliases.update({c.lower(): c for c in codes.values()})
aliases.update({'orissa':'OR','od':'OR','wb':'WB','jandk':'JK','andamanandnicobar':'AN'})
def states(text):
    clean = str(text).replace('&', ' and ')
    found = []
    for name, code in codes.items():
        pattern = r'\b' + r'\s+'.join(map(re.escape, name.split())) + r'\b'
        if re.search(pattern, clean, re.I) and code not in found: found.append(code)
    for match in re.finditer(r'\b(?:AP|TS|TN|KA|MH|MP|GJ|WB|BR|JH|UP|UK|HP|HR|PB|RJ|KL|AS|CG|OR|OD|JK)\b', clean):
        code = aliases.get(match.group().lower())
        if code and code not in found: found.append(code)
    return found
def blocks(filename, url):
    path = SOURCE / filename
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    body = E.fromstring(ZipFile(path).read('word/document.xml')).find('w:body', NS)
    output = []; table = 0
    for index, node in enumerate(body, 1):
        locator = {'url':url, 'document':filename, 'block':index, 'sha256':digest}
        if node.tag.endswith('}tbl'):
            table += 1
            rows = [['\n'.join(''.join(t.text or '' for t in p.findall('.//w:t', NS)) for p in cell.findall('w:p', NS)).strip() for cell in row.findall('w:tc', NS)] for row in node.findall('w:tr', NS)]
            output.append({'kind':'table', 'rows':rows, 'locator':{**locator,'table':table}})
        elif node.tag.endswith('}p'):
            text = ''.join(t.text or '' for t in node.findall('.//w:t', NS)).strip()
            if text: output.append({'kind':'paragraph','text':text,'locator':locator})
    return output
def exact_rows(table):
    for number, cells in enumerate(table['rows'][1:], 2):
        if any(cells): yield cells, {**table['locator'],'row':number}
def col(headers, pattern, default=0): return next((i for i,h in enumerate(headers) if re.search(pattern,h,re.I)),default)

cross = blocks('Cross Walks-2.0.docx', CROSSWALK)
context = blocks('production-context.docx', CONTEXT)
catalog = json.loads((SOURCE/'catalog-manifest.json').read_text())
drive = json.loads((SOURCE/'drive-assets.json').read_text())
drive_files = {row['slug']:row['files'] for row in drive}
notes = json.loads((SOURCE/'editor-notes.json').read_text())['values']
assets = read('asset-manifest.json'); source_records = read('source-records.json')
snapshot = read('procurement-records.json'); official = read('official-records.json')
summary = []
canonical = read('canonical-crops.json')
name_alias = {key(n): n for n in canonical}
name_alias.update({'chickpea':'Gram','gramchickpeachana':'Gram','rapeseedmustard':'Rapeseed & Mustard','soyabean':'Soybean','jute':'Jute','mesta':'Mesta','nigerseeds':'Niger Seeds','paddy':'Rice','greengram':'Moong','bengalgram':'Gram','arhar':'Pigeon Pea','tur':'Pigeon Pea','urad':'Black Gram'})
context_by_crop = {}; current = None
for block in context:
    # The supplied Production Context starts Barley prose without a crop heading.
    # Its exact first paragraph is block 2114; do not leak it into the Ragi section.
    if block['locator']['block']==2114:
        if 'Barley (Hordeum vulgare' not in block.get('text',''):raise ValueError('Barley context anchor changed')
        current='Barley';context_by_crop.setdefault(current,[])
    if block['kind']=='paragraph' and len(block['text'])<65:
        name = name_alias.get(key(block['text'].strip(' :–—-')))
        if name:
            current = name; context_by_crop.setdefault(name,[]); continue
    if current: context_by_crop[current].append(block)

# Jurisdiction is mapped only from the node heading or explicit facility fields,
# never from a catchment list. These three source-local names need state context.
jurisdiction_overrides = {('lentil',1):['MP'],('sesamum',0):['GJ'],('safflower',2):['MH']}
node_rules = []
for slug, (name, start, stop) in SECTIONS.items():
    section = [b for b in cross if start <= b['locator']['block'] < stop]
    if not section or section[0]['locator']['block']!=start: raise ValueError('Unresolved crop anchor '+slug)
    tables = [b for b in section if b['kind']=='table']
    calendar_start=next(b['locator']['block'] for b in section if b['kind']=='paragraph' and b['text'].startswith('State-Wise'))
    calendar_table=next(b for b in tables if b['locator']['block']>calendar_start)
    candidates=[b for b in tables if len(b['rows'][0])==4 and re.search('Belt',b['rows'][0][1],re.I)]
    if len(candidates)!=1:raise ValueError('Unresolved exact belt table: '+slug+' '+str([b['rows'][0] for b in tables]))
    belt_table=candidates[0]
    pest_table = next(b for b in tables if any('Target Pest' in h for h in b['rows'][0]))
    calendars = []; previous = ''; headers=calendar_table['rows'][0]
    state_col=col(headers,'^State'); variety_col=col(headers,'Cultivar|Variety|Selection',1); system_col=col(headers,'Season|Class|Primary Species|Ecosystem',2)
    window_cols=[i for i,h in enumerate(headers) if re.search('Window|Arrival|Harvest|Auction',h,re.I)]
    if not window_cols:raise ValueError('Source calendar has no named window: '+slug)
    for cells,loc in exact_rows(calendar_table):
        previous=cells[state_col] or previous
        calendars.append({'id':f"{slug}-calendar-{loc['table']}-{loc['row']}",'state':previous,'stateCodes':states(previous),'variety':cells[variety_col],'system':cells[system_col],'windows':[{'label':headers[i],'text':cells[i]} for i in window_cols if i<len(cells) and cells[i]],'cells':cells,'headers':headers,'locator':loc})
    belts=[];headers=belt_table['rows'][0]
    for cells,loc in exact_rows(belt_table):
        belts.append({'id':f"{slug}-belt-{loc['table']}-{loc['row']}",'origin':cells[1],'stateCodes':states(cells[1]),'variety':cells[0],'traits':cells[2],'destinations':cells[3],'cells':cells,'headers':headers,'locator':loc})
    marker=next(b['locator']['block'] for b in section if b['kind']=='paragraph' and b['text'].startswith('Processing Geographies'))
    node_starts=[b for b in section if b['kind']=='paragraph' and b['locator']['block']>marker and re.match(r'^\d+[.)]\s',b['text'])]
    # Each reviewed field-crop source contains three numbered processing nodes.
    if len(node_starts)!=3: raise ValueError('Expected three source processing nodes: '+slug)
    processing=[]
    for index,first in enumerate(node_starts):
        limit=node_starts[index+1]['locator']['block'] if index+1<len(node_starts) else first['locator']['block']+12
        paragraphs=[]
        for b in section:
            if b['kind']!='paragraph' or not first['locator']['block']<=b['locator']['block']<limit: continue
            if paragraphs and re.search(r'^(Detailed|Crop Phenology|Phenological|Phytosanitary|Germination|Seedling|Vegetative|State-Wise)',b['text'],re.I):break
            paragraphs.append(b)
        title=re.sub(r'^\d+[.)]\s*','',first['text'])
        jurisdiction=states(title)
        for b in paragraphs[1:]:
            if re.match(r'^(Key Facilities|Key Centers|Key Processors|Primary Facilities)',b['text'],re.I):jurisdiction=list(dict.fromkeys(jurisdiction+states(b['text'])))
        if (slug,index) in jurisdiction_overrides:jurisdiction=jurisdiction_overrides[slug,index]
        regional=[r for r in calendars if set(r['stateCodes'])&set(jurisdiction)]
        if 'basmati' in title.lower() and not re.search(r'non[- ]?basmati',title,re.I):
            regional=[r for r in regional if 'basmati' in (r['variety']+' '+r['system']).lower() and not re.search(r'non[- ]?basmati',r['variety']+' '+r['system'],re.I)]
        varieties=[{'variety':r['variety'],'geography':r['state'],'locator':r['locator'],'scope':'Regional calendar context; processor acceptance requires verification'} for r in regional]
        products=[b['text'] for b in paragraphs[1:] if re.match(r'^(Primary Products|Products|Strategic Role|Primary Role|Core Products)',b['text'],re.I)]
        catchment=[b['text'] for b in paragraphs[1:] if re.match(r'^(Catchment|Primary Supply|Supply Catchment)',b['text'],re.I)]
        processing.append({'id':f'{slug}-processing-{index}','name':title,'stateCodes':jurisdiction,'products':products,'catchment':catchment,'paragraphs':paragraphs,'locator':first['locator'],'index':index,'relatedVarieties':varieties})
        node_rules.append({'crop':slug,'node':index,'name':title,'states':jurisdiction,'source_block':first['locator']['block']})
    protection=[];locators=[]
    for cells,loc in exact_rows(pest_table):
        management=' '.join(cells[3:]);modes=[]
        if re.search(r'\b(?:EC|SC|WP|WDG|WG|SG|SL|SP|GR|fungicide|insecticide|fumigation|chemical)\b',management,re.I):modes.append('Chemical')
        if re.search(r'Trichoderma|Beauveria|Bacillus|predator|parasitoid|bio.?agent|biological|NPV',management,re.I):modes.append('Biological')
        if re.search(r'trap|hand.?pick|mechanical|destroy|remove.*infect|netting',management,re.I):modes.append('Mechanical')
        if re.search(r'rotation|sanitation|drain|spacing|resistan|avoid|early sow|clean|rogue|cultural',management,re.I):modes.append('Cultural')
        if re.search(r'quarantine|regulatory|MRL|export restriction',management,re.I):modes.append('Regulatory')
        if not modes:modes=['Source management context']
        # Formulations, rates, schedules and PHIs are deliberately not shipped.
        protection.append({'name':cells[1],'geography':cells[0],'impact':cells[2],'modes':modes,'type':'Diseases' if re.search('blight|wilt|rot|mildew|virus|rust|smut|blast',cells[1],re.I) else 'Pests','passage':f"{name} — source table {loc['table']}, row {loc['row']}. Management modes are descriptive; formulations, rates and schedules omitted."})
        locators.append({'name':cells[1],'locator':loc})
    context_sections=[context_by_crop.get(name,[])] if slug!='jute-mesta' else [context_by_crop.get('Jute',[]),context_by_crop.get('Mesta',[])]
    estimates=[];estimate_tables=[];national=[]
    for context_section in context_sections:
        current_states=[]
        for b in context_section:
            if b['kind']=='table':
                for cells,loc in exact_rows(b):estimate_tables.append({'id':f"{slug}-estimate-table-{loc['table']}-{loc['row']}",'stateCodes':states(cells[0]),'headers':b['rows'][0],'cells':cells,'locator':loc})
            else:
                value=b['text'];label=re.sub(r'^\d+[.)]\s*','',value).strip(' :')
                label=re.sub(r'\s*\([^)]*\)\s*$','',label).strip()
                if key(label) in aliases:current_states=[aliases[key(label)]]
                elif re.match(r'^\d+[.)]\s',value) and states(label):current_states=states(label)
                elif value.startswith('Industry Dynamics'):current_states=[]
                elif current_states and (re.search(r'\d.*%',value) or re.match('Remaining (District|Belt|Region)',value,re.I)):
                    estimates.append({'id':f"{slug}-estimate-{b['locator']['block']}",'stateCodes':current_states.copy(),'text':value,'locator':b['locator']})
                elif not national and re.search('produces|production|output',value,re.I):national.append(b)
    snapshot['crops'][slug]={'calendar':calendars,'belts':belts,'processing':processing,'protectionLocators':locators,'crosswalkUrl':CROSSWALK,'productionContext':{'summary':national,'rows':estimate_tables,'passages':estimates,'url':CONTEXT,'classification':'Production Context estimate; not official district APY'}}
    source_records[slug]={'protection':protection}
    crop_manifest=next(r for r in catalog['results'] if r['slug']==slug)
    screens={}
    for screen in ['002','003','006']:
        file=next(f for f in drive_files[slug] if f['title'].startswith('PR-'+screen) and f['mime_type']=='image/svg+xml')
        raw=(SOURCE/'originals'/f'{slug}-{screen}.svg').read_bytes()
        expected=next(f['sha256'] for f in crop_manifest['files'] if f['name']==file['title'])
        if hashlib.sha256(raw).hexdigest()!=expected:raise ValueError('Original SHA mismatch: '+slug+'/'+screen)
        screens[screen]={'footer':[], 'filename':file['title'],'sourceUrl':file['url'],'original':file['url'],'application':f'/gramshakti/crops/{slug}/{screen}.svg'}
    assets[slug]={'name':name,'screens':screens,'clusters':[{'name':r['name'],'texts':[str(i+1),r['name'],' / '.join(r['stateCodes']),'REGIONAL SOURCE VARIETIES',' · '.join(dict.fromkeys(v['variety'] for v in r['relatedVarieties'])) or 'No node-specific variety mapping supplied'],'color':['#EFA33A','#2E8B70','#8C6BB1'][i],'sourceContext':'\n'.join(b['text'] for b in r['paragraphs'])} for i,r in enumerate(processing)]}
    (ROOT/f'src/assets/gramshakti/{slug}').mkdir(parents=True,exist_ok=True)
    manifest=json.loads((SOURCE/'originals'/f'manifest-{slug}.json').read_text())
    (ROOT/f'src/assets/gramshakti/{slug}/manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    summary.append({'crop':name,'slug':slug,'calendar':len(calendars),'belts':len(belts),'processing':len(processing),'protection':len(protection),'estimate_passages':len(estimates),'editor_notes':next((row[1:4] for row in notes[1:] if row[0]==name),[])})

state_book=load_workbook(SOURCE/'statewide-production.xlsx',read_only=True,data_only=True)
state_sheet=next(s for s in state_book if s.title.startswith('Mandatory Crops Statewise'))
for slug in SECTIONS:official[slug]={'production':[]}
by_name={name:slug for slug,(name,*_) in SECTIONS.items()}
for number,row in enumerate(state_sheet.values,1):
    if len(row)<9 or row[2] not in by_name:continue
    if not isinstance(row[3],(int,float)):continue
    slug=by_name[row[2]]
    official[slug]['production'].append({'row':number,'year':row[0],'state':row[1],'tonnes':row[3],'authority':row[4],'publication':row[5],'url':row[6],'originalConversion':row[7],'status':row[8]})

district_book=load_workbook(SOURCE/'district-apy.xlsx',read_only=True,data_only=True)
registry={str(r[0]):{'title':r[1],'url':r[3]} for r in district_book['Data sources'].values if len(r)>4 and str(r[0]).startswith('SRC-')}
for slug in SECTIONS:snapshot['district'][slug]=[]
seen=set()
for number,row in enumerate(district_book['District APY Master'].values,1):
    if len(row)<23:continue
    native=name_alias.get(key(row[1]));slug=by_name.get(native)
    if native in ['Jute','Mesta']:slug='jute-mesta'
    if not slug:continue
    identity=(slug,row[3],row[2],row[0],native)
    if identity in seen:raise ValueError('Duplicate district observation '+str(identity))
    seen.add(identity);source=registry.get(str(row[18]),{})
    snapshot['district'][slug].append({'id':f'{slug}-district-{number}','cropLabel':str(row[1]),'row':number,'year':row[0],'district':row[2],'state':row[3],'area':row[4],'production':row[5],'yield':row[6],'yieldKind':'Derived yield' if row[18]=='SRC-08' and row[9]=='tonnes/ha' else 'Published/native yield','areaUnit':row[7],'productionUnit':row[8],'yieldUnit':row[9],'missingFlag':row[10] or 'NOT_SUPPLIED','normalizedProduction':row[11] if isinstance(row[11],(int,float)) else None,'derivedYield':row[12] if isinstance(row[12],(int,float)) else None,'sourceId':row[18],'areaPage':row[19],'productionPage':row[20],'yieldPage':row[21],'status':row[22],'sourceTitle':source.get('title','District APY Master'),'sourceUrl':source.get('url') if str(source.get('url','')).startswith('https://') else DISTRICT})
for row in summary:
    row.update(official_rows=len(official[row['slug']]['production']),district_rows=len(snapshot['district'][row['slug']]))
snapshot['snapshot']='2026-10-17'
write('asset-manifest.json',assets);write('source-records.json',source_records)
write('procurement-records.json',snapshot);write('official-records.json',official)
write('field-crops.json',list(SECTIONS))
write('field-crop-audit.json',{'snapshot':'2026-10-17','crosswalkUrl':CROSSWALK,'notesUrl':'https://docs.google.com/spreadsheets/d/1GZz93y1fusfbXxqElFZPahH8455yMYs_vtcLq1sdlLs/edit','crops':summary,'node_jurisdictions':node_rules,'blocked':[r for r in catalog['blockedCrops'] if r['crop'] in ['Maize','Sugarcane']]})
print(json.dumps(summary,ensure_ascii=False,indent=2))
